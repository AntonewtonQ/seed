"""Reproduce aggregate research outputs; raw and row-level data stay private.

Usage: .tools-venv/bin/python scripts/analyze-questionnaire.py --xlsx PRIVATE.xlsx
Never prints respondent records or contacts. Decisions apply to the recovered corpus.
"""
import argparse
import collections
import csv
import hashlib
import json
import os
import re
import unicodedata
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'docs/research'
PRIVATE = ROOT / '.local/research-2026-10-08'
DATE = '2026-10-08'
AGREE = ['Discordo totalmente', 'Discordo', 'Nem concordo nem discordo', 'Concordo', 'Concordo totalmente']
FREQ = ['Nunca', 'Raramente', 'Às vezes', 'Frequentemente', 'Quase sempre']
UTILITY = ['Nada útil', 'Pouco útil', 'Moderadamente útil', 'Muito útil', 'Extremamente útil']
TOOLS = ['E-mail', 'Ficheiros Excel/folhas de cálculo', 'Documentos partilhados (Google Drive, OneDrive, etc.)', 'Ferramentas genéricas de gestão de projecto (Trello, Asana, etc.)', 'Plataforma específica para gestão de investimento']
FEATURES = ['Gestão de perfil de startup (dados do negócio, sector, estágio)', 'Gestão de rondas de investimento (valores, condições, milestones)', 'Data room com controlo de acessos', 'Registo de contactos e reuniões com investidores', 'Registo de soft commits e estado das negociações', 'Painéis de indicadores (tempo, taxa de conversão, volume de soft commits)', 'Integração com ferramentas de assinatura electrónica (para term sheets)']
OPTIONS = {1:['Sim, aceito participar.', 'Não aceito participar.'], 2:['Fundador(a) ou membro de startup acelerada', 'Representante de aceleradora/incubadora', 'Investidor(a) (angel, venture capital, corporate venture)', 'Outra'], 3:['Menos de 1 ano','De 1 a 3 anos','De 4 a 6 anos','Mais de 6 anos'], 6:['Sim, existe um processo bem estruturado.','Existe um processo, mas pouco estruturado.','Não existe um processo claramente definido.','Não sei / Não se aplica.'],7:TOOLS,18:['1','2','3','4','5'],20:['Sim','Não'],26:FEATURES,27:['1','2','3','4','5'],28:['Sem interesse','Pouco interesse','Algum interesse','Muito interesse','Interesse máximo'],31:['Sim','Não']}
OPTIONS.update({i:AGREE for i in range(8,13)})
OPTIONS.update({i:FREQ for i in range(13,18)})
OPTIONS.update({i:UTILITY for i in range(21,26)})
IDS = ['timestamp','Q01','Q02','Q03','Q04','Q05','Q06','Q07',*[f'Q08{x}' for x in 'abcde'],*[f'Q09{x}' for x in 'abcde'],'Q10','Q11','Q12',*[f'Q13{x}' for x in 'abcde'],'Q14','Q15','Q16','Q17','Q18','Q19','Q20']
# Post-recovery content categories. Rules expose coding decisions, not participant text.
# Each category is binary per submission/question, regardless of repeated words.
RULES = {
19:[('P1','Atrasos ou dificuldade de converter interesse em compromisso','demor|tempos de decis|tempo de negoc|transformar o interesse|converter o interesse'),('P2','Burocracia e garantias','burocracia|garantias'),('P3','Comunicação, proposta de valor ou ausência de resposta','falta de resposta|comunicacao da proposta'),('P4','Acesso a rede de investidores','rede de investidores'),('P5','Demonstração de crescimento e métricas','crescimento sustentavel|metricas'),('P6','Ausência de fundos ou promessa não concretizada','nunca recebi|prometem'),('P7','Relação com investidores sem detalhe interpretável','mao que da')],
29:[('C1','Segurança','segur|encript'),('C2','Privacidade','privacidade'),('C3','Facilidade de uso e experiência','facil|intuitiv|experiencia do utilizador'),('C4','Legalidade e proteção de participações','compliance|participacoes|legais|legias|legalidade|conformidade com a lei'),('C5','Transparência, rastreabilidade e fidelidade dos dados','transparencia|rastreabilidade|dados fieis'),('C6','Suporte e continuidade','suporte|continuidade'),('C7','Operações financeiras sugeridas fora do âmbito','pagamento|transferencia|desembols|transacoes financeiras'),('C8','Colaboração entre fundadores','colaboracao'),('C9','Eficiência','eficiencia'),('C10','Acessibilidade','acessibilidade'),('C11','Certificações de segurança','certificacoes')],
30:[('F1','Apoio ou expectativa favorável','interessante|ideia muito bom|impulsionador|ajudara|funcional|escelentissima|nao parem'),('F2','Estratégia de captação','estrategia'),('F3','Acessibilidade, transparência e segurança','acessivel|transparente|segura'),('F4','Educação e networking','educacao|networking')]
}

def clean(v):
    if v is None:return ''
    if isinstance(v,float) and v.is_integer():return str(int(v))
    return str(v).strip()

def folded(v):
    return ''.join(c for c in unicodedata.normalize('NFKD',v.lower()) if not unicodedata.combining(c))

def split_multiple(value, options):
    """Match whole original options; commas inside options are never delimiters.

    Unmatched tail is one Other answer, including its commas. Unknown text before
    a later known option is ambiguous and fails for human review.
    """
    remaining=value; selected=[]
    while remaining:
        match=next((o for o in sorted(options,key=len,reverse=True) if remaining==o or remaining.startswith(o+', ')),None)
        if match:
            selected.append(match);remaining=remaining[len(match):].removeprefix(', ')
        else:
            if any(', '+o in remaining for o in options):raise ValueError('Ambiguous multiple response: private review required')
            selected.append(remaining);remaining=''
    if len(set(selected))!=len(selected):raise ValueError('Repeated selection: private review required')
    return selected

def csvout(name, rows, fields):
    path=OUT/name;path.parent.mkdir(parents=True,exist_ok=True)
    with path.open('w',newline='',encoding='utf-8') as f:
        w=csv.DictWriter(f,fieldnames=fields);w.writeheader();w.writerows(rows)

def main():
    ap=argparse.ArgumentParser();ap.add_argument('--xlsx',required=True);args=ap.parse_args()
    os.umask(0o077);PRIVATE.mkdir(parents=True,exist_ok=True);OUT.mkdir(parents=True,exist_ok=True)
    source=Path(args.xlsx);wb=openpyxl.load_workbook(source,data_only=False)
    sheets=[];records=[];headers=None
    for ws in wb.worksheets:
        allrows=list(ws.iter_rows(values_only=True));nonempty=[r for r in allrows if any(v is not None for v in r)]
        sheets.append({'sheet':ws.title,'state':ws.sheet_state,'physical_rows':ws.max_row,'nonempty_rows':len(nonempty),'columns':ws.max_column,'formula_cells':sum(c.data_type=='f' for row in ws for c in row)})
        if not nonempty:continue
        if headers is None:headers=[clean(v) for v in nonempty[0]]
        elif headers!=[clean(v) for v in nonempty[0]]:raise ValueError('Different sheet schema: review required')
        records.extend([list(r) for r in nonempty[1:]])
    assert len(headers)==33,'Changed questionnaire schema'
    normalized=[[clean(v) for v in r] for r in records]
    changes=[{'row':j+2,'column':IDS[i],'original':str(v),'normalized':normalized[j][i]} for j,r in enumerate(records) for i,v in enumerate(r) if v is not None and str(v)!=normalized[j][i]]
    (PRIVATE/'normalizations.json').write_text(json.dumps(changes,ensure_ascii=False,indent=2))
    # Inclusion is consent only; mixed profiles and multiterritorial responses retained.
    included=[r for r in normalized if r[1]=='Sim, aceito participar.']
    assert len(included)==24,'Changed corpus: reconsider coding and report narrative'
    n=len(included);contacts=[r[32].lower() for r in included if r[32]]
    contact_rows=[{'submission':i+1,'contact':r[32],'contact_permission':r[31]} for i,r in enumerate(included) if r[32]]
    (PRIVATE/'contacts.json').write_text(json.dumps(contact_rows,ensure_ascii=False,indent=2))
    # Research copy has neither timestamp nor contact; still private because of free text.
    (PRIVATE/'analysis-rows.json').write_text(json.dumps([dict(zip(IDS[1:32],r[1:32])) for r in included],ensure_ascii=False,indent=2))
    audit={'analysis_date':DATE,'xlsx_sha256':hashlib.sha256(source.read_bytes()).hexdigest(),'sheets':sheets,'exported_submissions':len(records),'consent_yes':n,'included_records':n,'excluded_records':len(records)-n,'unique_participants':'Não verificável: ausência de identificador autenticado; contactos distintos não comprovam pessoas distintas','contact_present':len(contacts),'distinct_contact_strings':len(set(contacts)),'contact_without_yes':sum(bool(r[32]) and r[31]!='Sim' for r in included),'exact_duplicate_excess':sum(v-1 for v in collections.Counter(map(tuple,normalized)).values()),'duplicate_answer_patterns_excluding_time_contact':sum(v-1 for v in collections.Counter(tuple(r[1:32]) for r in normalized).values()),'first_observed':min(r[0] for r in included),'last_observed':max(r[0] for r in included),'timezone':'Não indicado na exportação; não convertido','reference_screenshot_count':24,'reference_count_matches':n==24,'screenshots_available':False,'pdf_pages':240,'pdf_individual_blocks':24}
    (OUT/'audit.json').write_text(json.dumps(audit,ensure_ascii=False,indent=2))
    freq=[];dictionary=[];coding=[];qual=[];missing=[]
    for i,h in enumerate(headers):
        kind='nominal';order=OPTIONS.get(i,[])
        if i==0:kind='data/hora administrativa'
        elif i==32:kind='contacto pessoal separado'
        elif i in (7,26):kind='seleção múltipla'
        elif i in (19,29,30):kind='texto aberto'
        elif i in (3,18,27,28) or 8<=i<=17 or 21<=i<=25:kind='ordinal'
        elif i in (4,5):kind='texto geográfico; categoria composta preservada'
        blank=sum(not r[i] for r in included);missing.append({'variable':IDS[i],'answered':n-blank,'missing':blank,'base':n})
        dictionary.append({'variable':IDS[i],'xlsx_column':openpyxl.utils.get_column_letter(i+1),'question':h,'type':kind,'options_order':' | '.join(order),'required_pdf': 'não' if i in (0,5,19,30,31,32) else 'sim','pdf_page_first_block':{0:'não consta',1:1,2:2,3:2,4:2,5:2,6:3,7:3,**{k:4 for k in range(8,13)},**{k:5 for k in range(13,19)},19:6,20:6,**{k:7 for k in range(21,26)},26:8,27:8,28:9,29:9,30:9,31:9,32:10}[i],'treatment':'contacto não publicado; apenas presença' if i==32 else 'não publicar carimbo individual' if i==0 else 'trim; não alterar significado; ordinais numéricos convertidos para inteiro textual' })
        if i==0:continue
        if i==32:counts=collections.Counter('Preenchido' if r[i] else 'Ausente' for r in included);cats=['Preenchido','Ausente']
        elif i in (19,29,30):
            counts=collections.Counter('Ausente' if not r[i] else 'Sem conteúdo temático' if folded(r[i]) in ('nao','n/a','melhor nao') else 'Com conteúdo' for r in included);cats=['Com conteúdo','Sem conteúdo temático','Ausente']
            valid=counts['Com conteúdo'];theme_counts=collections.Counter()
            for j,r in enumerate(included,1):
                if not r[i] or folded(r[i]) in ('nao','n/a','melhor nao'):continue
                themes=[code for code,label,pattern in RULES[i] if re.search(pattern,folded(r[i]))]
                if not themes:raise ValueError('Uncoded text requires private review')
                coding.append({'submission':j,'question':IDS[i],'codes':themes})
                theme_counts.update(themes)
            for code,label,pattern in RULES[i]:qual.append({'question':IDS[i],'code':code,'theme':label,'n_records':theme_counts[code],'pct_all':round(100*theme_counts[code]/n,1),'base_all':n,'pct_content':round(100*theme_counts[code]/valid,1) if valid else None,'base_content':valid,'coding_rule_regex_normalized':pattern})
        elif i in (7,26):
            selections=[split_multiple(r[i],OPTIONS[i]) for r in included];counts=collections.Counter(o for s in selections for o in s);cats=OPTIONS[i]+sorted(set(counts)-set(OPTIONS[i]))
            # Arbitrary new Other strings require review before public aggregate release.
            allowed_other={'Nunca fiz isso','Participação em eventos de Startup, feiras'} if i==7 else set()
            assert set(counts)-set(OPTIONS[i])<=allowed_other
        else:
            vals=[('Outra' if i==2 and r[i] not in OPTIONS[2] else r[i]) or 'Ausente' for r in included]
            counts=collections.Counter(vals);cats=order+sorted(set(counts)-set(order))
        for cat in cats:
            count=counts[cat];freq.append({'variable':IDS[i],'category':cat,'n':count,'percent':round(100*count/n,1),'denominator':n,'missing':blank,'multiple':i in (7,26)})
        if i not in (19,29,30,32,7,26):assert sum(counts.values())==n
    csvout('tables/frequencies.csv',freq,list(freq[0]));csvout('variable-dictionary.csv',dictionary,list(dictionary[0]));csvout('tables/missingness.csv',missing,list(missing[0]));csvout('tables/qualitative.csv',qual,list(qual[0]))
    (PRIVATE/'coding-reviewed.json').write_text(json.dumps(coding,ensure_ascii=False,indent=2))
    # Subgroup comparison deliberately reports counts, no inferential test/ranking.
    groups=[]
    for profile in OPTIONS[2]:
        rr=[r for r in included if (r[2] if r[2] in OPTIONS[2] else 'Outra')==profile]
        groups.append({'profile':profile,'n':len(rr),'unstructured_or_weak_Q06':sum(r[6] in OPTIONS[6][1:3] for r in rr),'high_interest_Q16':sum(r[28] in OPTIONS[28][3:] for r in rr)})
    csvout('tables/subgroups.csv',groups,list(groups[0]))
    (OUT/'aggregate.json').write_text(json.dumps({'audit':audit,'frequencies':freq,'qualitative':qual,'subgroups':groups},ensure_ascii=False,indent=2))
    for p in OUT.rglob('*'):
        if p.is_file():p.chmod(0o644)
    print(f'Aggregate analysis: {n} included records; no respondent data printed.')

if __name__=='__main__':main()
